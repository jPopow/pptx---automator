#!/usr/bin/env python3
"""
Audience Book Updater - Direct PPTX Modification Tool

This script modifies a master PowerPoint template by injecting data from a CSV file
based on configuration settings, producing a populated presentation.
"""

import json
import os
import sys
from datetime import datetime
from pathlib import Path
import logging

import pandas as pd
from pptx import Presentation
from pptx.util import Pt
from pptx.chart.data import ChartData


# Configure logging
logging.basicConfig(
    level=logging.INFO,
    format='%(asctime)s - %(levelname)s - %(message)s',
    handlers=[
        logging.FileHandler('audience_book_updater.log'),
        logging.StreamHandler(sys.stdout)
    ]
)
logger = logging.getLogger(__name__)


class AudienceBookUpdater:
    """Main class for updating audience book presentations."""

    def __init__(self, working_dir='.'):
        """Initialize the updater with the working directory."""
        self.working_dir = Path(working_dir)
        self.template_path = self.working_dir / 'template.pptx'
        self.csv_path = self.working_dir / 'data.csv'
        self.config_path = self.working_dir / 'config.json'

        self.presentation = None
        self.data_df = None
        self.config = None

    def validate_input_files(self):
        """
        Phase 1: Validate that all required input files exist.

        Returns:
            bool: True if all files exist, False otherwise.
        """
        logger.info("Phase 1: Validating input files...")

        missing_files = []

        if not self.template_path.exists():
            missing_files.append(str(self.template_path))

        if not self.csv_path.exists():
            missing_files.append(str(self.csv_path))

        if not self.config_path.exists():
            missing_files.append(str(self.config_path))

        if missing_files:
            logger.error(f"Missing required files: {', '.join(missing_files)}")
            return False

        logger.info("All required files found.")
        return True

    def load_files(self):
        """
        Phase 1: Load all input files.

        Returns:
            bool: True if all files loaded successfully, False otherwise.
        """
        try:
            # Load CSV data
            logger.info(f"Loading CSV data from {self.csv_path}...")
            self.data_df = pd.read_csv(self.csv_path)
            logger.info(f"Loaded {len(self.data_df)} rows of data.")

            # Load config
            logger.info(f"Loading configuration from {self.config_path}...")
            with open(self.config_path, 'r', encoding='utf-8') as f:
                self.config = json.load(f)
            logger.info(f"Configuration loaded: {self.config.get('clientName', 'Unknown Client')}")

            # Load presentation
            logger.info(f"Loading PowerPoint template from {self.template_path}...")
            self.presentation = Presentation(str(self.template_path))
            logger.info(f"Template loaded with {len(self.presentation.slides)} slides.")

            return True

        except Exception as e:
            logger.error(f"Error loading files: {str(e)}")
            return False

    def find_shape_by_name(self, slide, shape_name):
        """
        Find a shape in a slide by its name.

        Args:
            slide: The slide to search in.
            shape_name: The name of the shape to find.

        Returns:
            The shape object if found, None otherwise.
        """
        for shape in slide.shapes:
            if shape.name == shape_name:
                return shape
        return None

    def find_shape_in_presentation(self, shape_name):
        """
        Find a shape across all slides in the presentation.

        Args:
            shape_name: The name of the shape to find.

        Returns:
            Tuple of (slide, shape) if found, (None, None) otherwise.
        """
        for slide in self.presentation.slides:
            shape = self.find_shape_by_name(slide, shape_name)
            if shape:
                return slide, shape
        return None, None

    def replace_text_in_shape(self, shape, new_text):
        """
        Replace text in a shape while preserving formatting.

        Args:
            shape: The shape containing the text to replace.
            new_text: The new text to insert.
        """
        if not hasattr(shape, 'text_frame'):
            logger.warning(f"Shape {shape.name} does not have a text frame.")
            return

        text_frame = shape.text_frame

        # Clear existing paragraphs except the first one
        for _ in range(len(text_frame.paragraphs) - 1):
            text_frame.paragraphs[1]._element.getparent().remove(text_frame.paragraphs[1]._element)

        # Update the first paragraph
        if text_frame.paragraphs:
            paragraph = text_frame.paragraphs[0]
            if paragraph.runs:
                # Preserve formatting from first run
                first_run = paragraph.runs[0]
                font_name = first_run.font.name
                font_size = first_run.font.size
                font_bold = first_run.font.bold
                font_italic = first_run.font.italic

                # Clear all runs
                for _ in range(len(paragraph.runs)):
                    paragraph.runs[0]._element.getparent().remove(paragraph.runs[0]._element)

                # Add new text with preserved formatting
                new_run = paragraph.add_run()
                new_run.text = new_text
                if font_name:
                    new_run.font.name = font_name
                if font_size:
                    new_run.font.size = font_size
                if font_bold is not None:
                    new_run.font.bold = font_bold
                if font_italic is not None:
                    new_run.font.italic = font_italic
            else:
                # No existing runs, just set text
                paragraph.text = new_text

    def populate_text_fields(self):
        """
        Phase 2: Populate text fields based on configuration.

        Returns:
            bool: True if successful, False otherwise.
        """
        logger.info("Phase 2: Populating text fields...")

        try:
            personas = self.config.get('personas', [])

            for persona in personas:
                template_slot = persona.get('templateSlot')
                qualitative_text = persona.get('qualitativeText', {})

                logger.info(f"Processing persona slot {template_slot}...")

                for field_name, text_value in qualitative_text.items():
                    # Construct placeholder name
                    placeholder_name = f"P{template_slot}_TEXT::{field_name}"

                    # Find the shape
                    slide, shape = self.find_shape_in_presentation(placeholder_name)

                    if shape:
                        logger.info(f"  Updating {placeholder_name}")
                        self.replace_text_in_shape(shape, text_value)
                    else:
                        logger.warning(f"  Shape not found: {placeholder_name}")

            logger.info("Text population completed.")
            return True

        except Exception as e:
            logger.error(f"Error populating text fields: {str(e)}")
            return False

    def get_chart_data_from_csv(self, category, subcategories, audience_name):
        """
        Extract chart data from the CSV for a specific audience.

        Args:
            category: The main category to filter by.
            subcategories: List of subcategory dictionaries with csvName and displayName.
            audience_name: The audience name to get data for.

        Returns:
            Dictionary with categories and their values {displayName: (percent, index)}.
        """
        chart_data = {}

        # Filter data by category
        category_data = self.data_df[self.data_df['Category'] == category]

        for subcat in subcategories:
            csv_name = subcat['csvName']
            display_name = subcat['displayName']

            # Find the row for this subcategory
            subcat_row = category_data[category_data['Subcategory'] == csv_name]

            if not subcat_row.empty:
                # Get the percentage and index columns for this audience
                percent_col = f"{audience_name}_Audience %"
                index_col = f"{audience_name}_Index"

                if percent_col in subcat_row.columns and index_col in subcat_row.columns:
                    percent = subcat_row[percent_col].values[0]
                    index = subcat_row[index_col].values[0]
                    chart_data[display_name] = (percent, index)
                else:
                    logger.warning(f"Columns not found for audience {audience_name}")
            else:
                logger.warning(f"Subcategory {csv_name} not found in category {category}")

        return chart_data

    def update_chart_data(self, chart_shape, chart_data_dict):
        """
        Update a chart with new data.

        Args:
            chart_shape: The shape containing the chart.
            chart_data_dict: Dictionary with categories and values.
        """
        if not chart_shape.has_chart:
            logger.warning(f"Shape {chart_shape.name} does not contain a chart.")
            return

        chart = chart_shape.chart

        # Create new chart data
        chart_data = ChartData()
        chart_data.categories = list(chart_data_dict.keys())

        # Add series for percentage and index
        percentages = [chart_data_dict[cat][0] for cat in chart_data.categories]
        indices = [chart_data_dict[cat][1] for cat in chart_data.categories]

        # Assuming the chart has two series: Audience % and Index
        # We'll replace the data for both series
        chart_data.add_series('Audience %', percentages)
        chart_data.add_series('Index', indices)

        # Replace the chart data
        chart.replace_data(chart_data)

    def populate_charts(self):
        """
        Phase 3: Populate charts based on configuration.

        Returns:
            bool: True if successful, False otherwise.
        """
        logger.info("Phase 3: Populating charts...")

        try:
            charts_config = self.config.get('charts', [])
            personas = self.config.get('personas', [])

            for chart_config in charts_config:
                chart_slot = chart_config.get('chartSlot')
                chart_title = chart_config.get('title')
                category = chart_config.get('category')
                subcategories = chart_config.get('subcategories', [])

                logger.info(f"Processing chart slot {chart_slot}: {chart_title}...")

                # Update chart for each persona
                for persona in personas:
                    template_slot = persona.get('templateSlot')
                    csv_audience_name = persona.get('csvAudienceName')

                    # Construct chart shape name
                    chart_shape_name = f"P{template_slot}_CHART::{chart_slot}"

                    # Find the chart shape
                    slide, chart_shape = self.find_shape_in_presentation(chart_shape_name)

                    if chart_shape:
                        logger.info(f"  Updating chart {chart_shape_name} for {csv_audience_name}")

                        # Get data from CSV
                        chart_data = self.get_chart_data_from_csv(
                            category, subcategories, csv_audience_name
                        )

                        # Update the chart
                        if chart_data:
                            self.update_chart_data(chart_shape, chart_data)
                        else:
                            logger.warning(f"  No data found for chart {chart_shape_name}")
                    else:
                        logger.warning(f"  Chart shape not found: {chart_shape_name}")

                    # Update chart title
                    title_shape_name = f"P{template_slot}_TEXT::Chart{chart_slot}_Title"
                    slide, title_shape = self.find_shape_in_presentation(title_shape_name)

                    if title_shape:
                        logger.info(f"  Updating chart title {title_shape_name}")
                        self.replace_text_in_shape(title_shape, chart_title)
                    else:
                        logger.warning(f"  Chart title shape not found: {title_shape_name}")

            logger.info("Chart population completed.")
            return True

        except Exception as e:
            logger.error(f"Error populating charts: {str(e)}")
            return False

    def get_configured_slots(self):
        """
        Get the persona and chart slots that are configured.

        Returns:
            Tuple of (persona_slots, chart_slots).
        """
        persona_slots = set()
        chart_slots = set()

        for persona in self.config.get('personas', []):
            persona_slots.add(persona.get('templateSlot'))

        for chart in self.config.get('charts', []):
            chart_slots.add(chart.get('chartSlot'))

        return persona_slots, chart_slots

    def cleanup_unused_slides(self, configured_persona_slots):
        """
        Phase 4: Remove slides for unused persona slots.

        Args:
            configured_persona_slots: Set of persona slots that are configured.

        Returns:
            bool: True if successful, False otherwise.
        """
        logger.info("Phase 4: Cleaning up unused slides...")

        try:
            # Determine the maximum possible persona slot
            # We'll check for patterns like P1_, P2_, etc.
            max_slot = 10  # Reasonable maximum

            unused_slots = set(range(1, max_slot + 1)) - configured_persona_slots

            if not unused_slots:
                logger.info("No unused persona slots found.")
                return True

            logger.info(f"Unused persona slots: {sorted(unused_slots)}")

            # Iterate through slides in reverse to avoid index shifting
            slides_to_remove = []

            for idx in range(len(self.presentation.slides) - 1, -1, -1):
                slide = self.presentation.slides[idx]

                # Check if slide contains shapes for unused personas
                for unused_slot in unused_slots:
                    prefix = f"P{unused_slot}_"

                    # Check if any shape in the slide starts with this prefix
                    has_unused_persona = any(
                        shape.name.startswith(prefix)
                        for shape in slide.shapes
                        if hasattr(shape, 'name')
                    )

                    if has_unused_persona:
                        logger.info(f"  Removing slide {idx + 1} (contains {prefix}...)")
                        slides_to_remove.append(idx)
                        break  # No need to check other unused slots for this slide

            # Remove slides
            for slide_idx in slides_to_remove:
                rId = self.presentation.slides._sldIdLst[slide_idx].rId
                self.presentation.part.drop_rel(rId)
                del self.presentation.slides._sldIdLst[slide_idx]

            logger.info(f"Removed {len(slides_to_remove)} unused slides.")
            return True

        except Exception as e:
            logger.error(f"Error cleaning up unused slides: {str(e)}")
            return False

    def cleanup_unused_charts(self, configured_persona_slots, configured_chart_slots):
        """
        Phase 4: Remove unused chart shapes from slides.

        Args:
            configured_persona_slots: Set of persona slots that are configured.
            configured_chart_slots: Set of chart slots that are configured.

        Returns:
            bool: True if successful, False otherwise.
        """
        logger.info("Cleaning up unused chart shapes...")

        try:
            # Determine max chart slot
            max_chart_slot = 10  # Reasonable maximum

            unused_chart_slots = set(range(1, max_chart_slot + 1)) - configured_chart_slots

            if not unused_chart_slots:
                logger.info("No unused chart slots found.")
                return True

            logger.info(f"Unused chart slots: {sorted(unused_chart_slots)}")

            # Iterate through all slides
            for slide_idx, slide in enumerate(self.presentation.slides):
                shapes_to_remove = []

                # Check each configured persona's unused charts
                for persona_slot in configured_persona_slots:
                    for chart_slot in unused_chart_slots:
                        # Chart shape name
                        chart_name = f"P{persona_slot}_CHART::{chart_slot}"
                        # Chart title name
                        title_name = f"P{persona_slot}_TEXT::Chart{chart_slot}_Title"

                        # Find and mark for removal
                        for shape in slide.shapes:
                            if hasattr(shape, 'name'):
                                if shape.name == chart_name or shape.name == title_name:
                                    shapes_to_remove.append(shape)

                # Remove shapes
                for shape in shapes_to_remove:
                    logger.info(f"  Removing shape {shape.name} from slide {slide_idx + 1}")
                    sp = shape.element
                    sp.getparent().remove(sp)

            logger.info("Chart cleanup completed.")
            return True

        except Exception as e:
            logger.error(f"Error cleaning up unused charts: {str(e)}")
            return False

    def save_presentation(self):
        """
        Phase 5: Save the modified presentation.

        Returns:
            str: The output filename if successful, None otherwise.
        """
        logger.info("Phase 5: Saving presentation...")

        try:
            # Construct output filename
            client_name = self.config.get('clientName', 'Client')
            date_str = datetime.now().strftime('%Y%m%d')
            output_filename = f"{client_name} - Audience Book - {date_str}.pptx"
            output_path = self.working_dir / output_filename

            # Save the presentation
            self.presentation.save(str(output_path))

            logger.info(f"Presentation saved: {output_filename}")
            return output_filename

        except Exception as e:
            logger.error(f"Error saving presentation: {str(e)}")
            return None

    def run(self):
        """
        Main execution method that runs all phases.

        Returns:
            bool: True if successful, False otherwise.
        """
        logger.info("=" * 60)
        logger.info("AUDIENCE BOOK UPDATER - Starting")
        logger.info("=" * 60)

        # Phase 1: Validation and Loading
        if not self.validate_input_files():
            logger.error("Input file validation failed. Exiting.")
            return False

        if not self.load_files():
            logger.error("File loading failed. Exiting.")
            return False

        # Phase 2: Text Population
        if not self.populate_text_fields():
            logger.error("Text population failed. Exiting.")
            return False

        # Phase 3: Chart Population
        if not self.populate_charts():
            logger.error("Chart population failed. Exiting.")
            return False

        # Phase 4: Dynamic Cleanup
        persona_slots, chart_slots = self.get_configured_slots()

        if not self.cleanup_unused_slides(persona_slots):
            logger.error("Slide cleanup failed. Exiting.")
            return False

        if not self.cleanup_unused_charts(persona_slots, chart_slots):
            logger.error("Chart cleanup failed. Exiting.")
            return False

        # Phase 5: Finalization
        output_file = self.save_presentation()

        if output_file:
            logger.info("=" * 60)
            logger.info("SUCCESS! Audience Book has been created.")
            logger.info(f"Output file: {output_file}")
            logger.info("=" * 60)
            return True
        else:
            logger.error("Failed to save presentation. Exiting.")
            return False


def main():
    """Main entry point for the script."""
    print("\n" + "=" * 60)
    print("AUDIENCE BOOK UPDATER")
    print("Direct PPTX Modification Tool")
    print("=" * 60 + "\n")

    # Create and run the updater
    updater = AudienceBookUpdater()

    success = updater.run()

    if success:
        print("\n✓ Process completed successfully!")
        sys.exit(0)
    else:
        print("\n✗ Process failed. Check the log file for details.")
        sys.exit(1)


if __name__ == "__main__":
    main()
